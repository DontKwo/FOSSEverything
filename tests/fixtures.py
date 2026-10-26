import hashlib
import json
import math
import struct
import subprocess
import wave
import zipfile
from pathlib import Path

import docx
import imageio_ffmpeg
from PIL import Image, ImageDraw
import pymupdf
import qrcode
import trimesh
import yaml


def calc_sample_pdf(path: Path) -> str:
    # Generate 3-page text PDF.
    doc = pymupdf.open()
    for i in range(1, 4):
        page = doc.new_page(width=612, height=792)
        page.insert_text((72, 72), f"Sample PDF Document - Page {i}", fontsize=18)
        page.insert_text((72, 120), f"Body text content on page {i} for testing extraction and OCR.", fontsize=11)
    doc.save(str(path))
    doc.close()
    return str(path.resolve())


def calc_scan_pdf(path: Path) -> str:
    # Generate 1-page scanned image PDF.
    img = Image.new("RGB", (600, 800), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((50, 50), "Scanned document text for optical character recognition testing.", fill=(0, 0, 0))
    img_path = path.with_suffix(".temp.png")
    img.save(str(img_path))

    doc = pymupdf.open()
    page = doc.new_page(width=600, height=800)
    page.insert_image(page.rect, filename=str(img_path))
    doc.save(str(path))
    doc.close()
    if img_path.exists():
        img_path.unlink()
    return str(path.resolve())


def calc_sample_images(path1: Path, path2: Path, photo_path: Path) -> tuple[str, str, str]:
    # Generate test images with distinct shapes and colors.
    img1 = Image.new("RGBA", (800, 600), color=(30, 45, 75, 255))
    draw1 = ImageDraw.Draw(img1)
    draw1.rectangle([50, 50, 200, 200], fill=(220, 50, 50, 255))
    draw1.ellipse([250, 50, 400, 200], fill=(50, 200, 50, 255))
    draw1.text((50, 250), "Sample Test Pattern 1", fill=(240, 240, 240, 255))
    img1.save(str(path1))

    img2 = Image.new("RGBA", (640, 480), color=(50, 60, 80, 255))
    draw2 = ImageDraw.Draw(img2)
    draw2.rectangle([40, 40, 180, 180], fill=(70, 130, 240, 255))
    draw2.text((40, 220), "Sample Test Pattern 2", fill=(255, 255, 255, 255))
    img2.save(str(path2))

    photo = Image.new("RGB", (800, 600), color=(120, 140, 160))
    draw_photo = ImageDraw.Draw(photo)
    draw_photo.rectangle([100, 100, 700, 500], fill=(200, 180, 150))
    photo.save(str(photo_path), quality=90)
    return str(path1.resolve()), str(path2.resolve()), str(photo_path.resolve())


def calc_tone_wav(path: Path) -> str:
    # Generate 1-second 440 Hz sine wave audio file.
    sample_rate = 44100
    duration_s = 1.0
    freq_hz = 440.0
    num_samples = int(sample_rate * duration_s)

    with wave.open(str(path), "w") as wav_file:
        wav_file.setnchannels(1)
        wav_file.setsampwidth(2)
        wav_file.setframerate(sample_rate)
        for i in range(num_samples):
            val = int(32767.0 * 0.5 * math.sin(2.0 * math.pi * freq_hz * (i / sample_rate)))
            wav_file.writeframes(struct.pack("<h", val))
    return str(path.resolve())


def calc_clip_mp4(path: Path) -> str:
    # Generate 1-second synthetic MP4 video clip with audio stream.
    ffmpeg_exe = imageio_ffmpeg.get_ffmpeg_exe()
    cmd = [
        ffmpeg_exe,
        "-y",
        "-f",
        "lavfi",
        "-i",
        "testsrc=duration=1:size=160x120:rate=10",
        "-f",
        "lavfi",
        "-i",
        "sine=frequency=1000:duration=1",
        "-c:v",
        "libx264",
        "-c:a",
        "aac",
        "-pix_fmt",
        "yuv420p",
        "-shortest",
        str(path),
    ]
    subprocess.run(cmd, capture_output=True, check=True)
    return str(path.resolve())


def calc_cube_stl(path: Path) -> str:
    # Generate 3D cube mesh in STL format.
    mesh = trimesh.creation.box(extents=(10.0, 10.0, 10.0))
    mesh.export(str(path))
    return str(path.resolve())


def calc_sample_docx(path: Path) -> str:
    # Generate sample DOCX file with headings and paragraphs.
    doc = docx.Document()
    doc.add_heading("Sample Document Title", level=1)
    doc.add_paragraph("First paragraph of test document.")
    doc.add_paragraph("Second paragraph with additional sample text.")
    doc.save(str(path))
    return str(path.resolve())


def calc_qr_png(path: Path) -> str:
    # Generate QR code image.
    img = qrcode.make("https://example.com/foss-test")
    img.save(str(path))
    return str(path.resolve())


def make_fixtures(target_dir: Path) -> dict[str, str]:
    # Create test fixture files in given directory and return path map.
    target_dir.mkdir(parents=True, exist_ok=True)
    fixtures: dict[str, str] = {}

    pdf_path = target_dir / "sample.pdf"
    fixtures["sample.pdf"] = calc_sample_pdf(pdf_path)

    scan_path = target_dir / "scan.pdf"
    fixtures["scan.pdf"] = calc_scan_pdf(scan_path)

    p1, p2, photo = calc_sample_images(target_dir / "sample.png", target_dir / "sample2.png", target_dir / "photo.jpg")
    fixtures["sample.png"] = p1
    fixtures["sample2.png"] = p2
    fixtures["photo.jpg"] = photo

    tone_path = target_dir / "tone.wav"
    fixtures["tone.wav"] = calc_tone_wav(tone_path)

    clip_path = target_dir / "clip.mp4"
    fixtures["clip.mp4"] = calc_clip_mp4(clip_path)

    txt_path = target_dir / "sample.txt"
    txt_path.write_text("Hello FOSSEverything.\nSecond line of plain text.\n", encoding="utf-8")
    fixtures["sample.txt"] = str(txt_path.resolve())

    data_json_path = target_dir / "data.json"
    data_json_path.write_text(json.dumps({"name": "Test", "items": [1, 2, 3], "active": True}, indent=2), encoding="utf-8")
    fixtures["data.json"] = str(data_json_path.resolve())

    data_yaml_path = target_dir / "data.yaml"
    data_yaml_path.write_text(yaml.dump({"project": "FOSSEverything", "version": 1.0}), encoding="utf-8")
    fixtures["data.yaml"] = str(data_yaml_path.resolve())

    data_csv_path = target_dir / "data.csv"
    data_csv_path.write_text("id,name,value\n1,alpha,10.5\n2,beta,20.0\n", encoding="utf-8")
    fixtures["data.csv"] = str(data_csv_path.resolve())

    md_path = target_dir / "sample.md"
    md_path.write_text("# Heading\n\n- item 1\n- item 2\n\nParagraph text.\n", encoding="utf-8")
    fixtures["sample.md"] = str(md_path.resolve())

    zip_path = target_dir / "sample.zip"
    with zipfile.ZipFile(str(zip_path), "w") as zf:
        zf.writestr("inner.txt", "Inner file contents.")
    fixtures["sample.zip"] = str(zip_path.resolve())

    stl_path = target_dir / "cube.stl"
    fixtures["cube.stl"] = calc_cube_stl(stl_path)

    docx_path = target_dir / "sample.docx"
    fixtures["sample.docx"] = calc_sample_docx(docx_path)

    qr_path = target_dir / "qr.png"
    fixtures["qr.png"] = calc_qr_png(qr_path)

    manifest_path = target_dir / "manifest.sha256"
    sha = hashlib.sha256(b"Hello FOSSEverything.\nSecond line of plain text.\n").hexdigest()
    manifest_path.write_text(f"{sha}  sample.txt\n", encoding="utf-8")
    fixtures["manifest.sha256"] = str(manifest_path.resolve())

    return fixtures

